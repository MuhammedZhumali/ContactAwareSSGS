"""Focused local validation of revision 2; does not execute policy searches."""
import ast,collections,hashlib,json,re,subprocess,sys,zipfile
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'target/aitu-revision-v2';DOC=ROOT/'docs/aitu-revision-v2'
def load(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def sha(p):
 with p.open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()
def save(p,x):p.write_text(json.dumps(x,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')

def main():
 preserved=load(BASE/'preservation.json')
 for p,h in preserved.items():assert sha(ROOT/p)==h,p
 frozen=load(ROOT/'target/final-research/final-v2-sources.json')
 sources={p:sha(ROOT/p)==h for p,h in frozen.items() if p.startswith('src/main/')}
 assert len(sources)==112 and all(sources.values())
 protocol=load(BASE/'protocol.json');registration=load(BASE/'controller-registration.json')
 assert sha(ROOT/'scripts/run_aitu_v2.py')==protocol['runnerSha256']
 assert sha(DOC/'protocol.md')==protocol['protocolSha256']
 for p,h in protocol['inputs'].items():assert sha(ROOT/p)==h,p
 assert sha(BASE/'harness/manifest.json')==protocol['harnessManifestSha256']
 assert sha(BASE/'protocol.json')==registration['sourceProtocolSha256']
 assert sha(BASE/'root-summary.json')==registration['rootSummarySha256']
 assert (protocol['searchSeconds'],protocol['decisionSeconds'],protocol['controllerSeconds'],protocol['controllerWatchdogSeconds'])==(30,32,180,210)
 audit=load(BASE/'audit.json');roots=load(BASE/'root-tables.json');runs=load(BASE/'controller-tables.json')
 assert len(roots)==len(protocol['cases'])==6 and len(runs)==len(registration['cases'])==4
 assert collections.Counter(r['status'] for r in roots)=={'COMPLETE':2,'INCOMPLETE':4}
 assert collections.Counter(r['status'] for r in runs)=={'FINISHED':3,'CONTROLLER_TIME_LIMIT':1}
 assert all(r['replayAccepted'] and r['retainedScore'] is None for r in roots)
 assert sum(r['certificate'] for r in roots)==2
 for rule in ['CANONICAL','CONTINUE_WAIT']:
  assert registration['triggers'][rule]==min(r['nodes'] for r in roots if r['rule']==rule and r['beatsOldPractical'])==150000
 for r in runs:
  if r['status']!='FINISHED':assert r['finalScore'] is None and r['prefixScore'] is not None
 for p,h in audit['inputHashes'].items():assert sha(ROOT/p)==h,p
 desc=load(BASE/'input-descriptors.json');assert len(desc['models'])==11
 for model in desc['models']:assert sha(ROOT/model['path'])==model['sha256']
 for p,h in desc['codeSourceHashesMatchFrozen'].items():assert sha(ROOT/p)==h==frozen[p]
 rebuild=load(BASE/'rebuild-check/validation.json')
 assert rebuild['status']=='PASS' and len(rebuild['newClasses'])==1 and len(rebuild['priorClasses'])==26
 assert all(rebuild['newClasses'].values()) and all(rebuild['priorClasses'].values())
 article=load(DOC/'article-content.json');refs=[]
 for b in article['blocks']:
  for n in re.findall(r'\[(\d+)\]',b.get('text','')):
   if n not in refs:refs.append(n)
 assert refs==list(map(str,range(1,21))) and len(article['references'])==20
 assert len(article['abstract'].split())==269 and len(article['abstract'])==1672 and len(article['title'])==78
 build=load(BASE/'document-build.json');visual=load(DOC/'visual-review.json')
 assert visual['status']=='PASS'
 checks={}
 for stem,tex in [('article','main.tex'),('supplement','supplement.tex')]:
  p=DOC/(stem+'.docx');d=Document(p);s=d.sections[0]
  assert len(d.sections)==1
  for actual,want in [(s.page_width.cm,21),(s.page_height.cm,29.7),(s.top_margin.cm,2),(s.bottom_margin.cm,2.5),(s.left_margin.cm,3),(s.right_margin.cm,1.5)]:assert abs(actual-want)<.01
  pars=list(d.paragraphs)+[p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs]
  for paragraph in pars:
   for run in paragraph.runs:
    if run.text:assert run.font.name=='Times New Roman' and run.font.size.pt==12,(stem,run.text)
  for t in d.tables:
   assert t._tbl.tblPr.find(qn('w:tblLayout')).get(qn('w:type'))=='fixed'
   assert t.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None
  with zipfile.ZipFile(p) as z,zipfile.ZipFile(ROOT/'target/aitu-revision/reference-materials/aitu-article-template.docx') as template:
   for part in build['preserveOnlyParts']:assert z.read(part)==template.read(part)
   omml=z.read('word/document.xml').count(b'<m:oMath>');assert omml==(3 if stem=='article' else 0)
  assert sha(p)==build['files'][stem]['docxSha256']==visual['documents'][stem]['docxSha256']
  tp=DOC/'article-tex'/tex
  assert sha(tp)==build['files'][stem]['texSha256']==visual['latex'][stem]['sha256']
  assert visual['latex'][stem]['status']=='success'
  assert not re.search(r'\\url\{[^}]*[.,;:]\}',tp.read_text(encoding='utf-8'))
  checks[stem]={'sha256':sha(p),'pages':visual['documents'][stem]['pages'],'nativeOMML':omml,'tables':len(d.tables),'font':'Times New Roman 12pt','geometry':'A4; margins top2 bottom2.5 left3 right1.5cm','preservedTemplateParts':len(build['preserveOnlyParts'])}
 scripts=list((ROOT/'scripts').glob('*aitu*v2*.py'))
 for p in scripts:ast.parse(p.read_text(encoding='utf-8-sig'),filename=str(p))
 diff=subprocess.run(['git','diff','--check'],cwd=ROOT,capture_output=True,text=True)
 assert diff.returncode==0,diff.stdout+diff.stderr
 value={'status':'PASS','roots':audit['roots'],'controllers':audit['controllers'],'historicalFilesUnchanged':len(preserved),'frozenProductionSourcesMatch':len(sources),'auditedRawHashes':len(audit['inputHashes']),'frozenInputConfigurations':11,'recompiledClassesByteIdentical':27,'article':{'references':20,'firstCitationOrder':refs,'abstractWords':269,'abstractCharacters':1672,'titleCharacters':78},'documents':checks,'newPythonScriptsParsed':len(scripts),'gitDiffCheck':'PASS','boundary':'Local preservation, raw result checks, bytecode rebuild, document structure and recorded page/TeX QA. No new scientific runs or Maven suite.'}
 save(DOC/'validation.json',value);save(BASE/'final-validation.json',value)
 save(BASE/'final-environment.json',{'identity':load(BASE/'identity.json'),'finalGitCommit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),'finalGitStatus':subprocess.check_output(['git','status','--porcelain=v1'],cwd=ROOT,text=True),'python':sys.version,'hardware':load(ROOT/'target/aitu-revision/environment-confirmed.json'),'jarSha256':sha(ROOT/'target/final-research/final-v2.jar'),'boundary':'Dirty tree; frozen source and JAR hashes, rather than HEAD alone, identify the executable.'})
 print(json.dumps({k:v for k,v in value.items() if k not in ['article','documents']},indent=2))

if __name__=='__main__':main()
