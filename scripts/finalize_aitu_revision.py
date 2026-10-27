"""Final evidence validation and compact exports; never runs policy experiments."""
import json,csv,collections,hashlib,subprocess,sys,shutil,datetime,zipfile
from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[1];BASE=ROOT/'target/aitu-revision';DOC=ROOT/'docs/aitu-revision';TABLES=DOC/'tables';TABLES.mkdir(exist_ok=True)
def sha(p):
 with open(p,'rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()
def load(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def save(p,v):p.write_text(json.dumps(v,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
def csvwrite(p,rows):
 fields=list(dict.fromkeys(k for x in rows for k in x))
 with p.open('w',encoding='utf-8-sig',newline='') as f:w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)
allrows=[];stages={};hashed=0
for stage in ['stage1','stage2','stage3','stage4']:
 d=load(BASE/stage/'audit.json');stages[stage]=d['denominator'];assert d['denominator']['notRun']==0
 for name,h in d['inputHashes'].items():assert sha(ROOT/name)==h,name;hashed+=1
 for ext in ['json','csv']:shutil.copy2(BASE/stage/f'tables.{ext}',TABLES/f'{stage}.{ext}')
 allrows += [dict(stage=stage,**r) for r in load(BASE/stage/'tables.json')]
for ext in ['json','csv']:shutil.copy2(BASE/'stage1'/f'root-tables.{ext}',TABLES/f'root-diagnostics.{ext}')
csvwrite(TABLES/'all-controller-cells.csv',allrows)
new=[r for r in allrows if not r['reused']];reused=[r for r in allrows if r['reused']]
assert len(new)==110 and len(reused)==6
assert collections.Counter(r['status'] for r in new)=={'FINISHED':105,'SEARCH_INCOMPLETE':5}
assert sum(len(load(BASE/stage/'protocol.json')['cases']) for stage in stages)==122
readiness=[];conditions=[]
for case,result in load(BASE/'stage4/audit.json')['runs'].items():
 for i in result['recovery']['readinessIntervals']:
  readiness.append({'case':case,**{k:v for k,v in i.items() if k not in ['keyIds','originalDeadlines']},'keyIds':json.dumps(i['keyIds']),'originalDeadlines':json.dumps(i['originalDeadlines'])})
 for i in result['recovery']['conditionIntervals']:conditions.append({'case':case,**{k:v for k,v in i.items() if k!='conditions'},**i['conditions']})
csvwrite(TABLES/'recovery-readiness.csv',readiness);csvwrite(TABLES/'overlapping-wait-conditions.csv',conditions)
preserved=load(BASE/'preservation.json')
for name,h in preserved.items():assert sha(ROOT/name)==h,name
frozen=load(ROOT/'target/final-research/final-v2-sources.json');sourcechecks={p:sha(ROOT/p)==h for p,h in frozen.items() if p.startswith('src/main/')}
assert all(sourcechecks.values())
rebuild=load(BASE/'rebuild-check/validation.json');assert rebuild['status']=='PASS' and all(rebuild['checks'].values())
article=load(DOC/'article-content.json');assert len(article['references'])==20 and 250<=len(article['abstract'].split())<=280 and len(article['abstract'])<=1900
assert len(article['title'])<=96
refOrder=[]
import re
for block in article['blocks']:
 for n in re.findall(r'\[(\d+)\]',block.get('text','')):
  if n not in refOrder:refOrder.append(n)
assert refOrder==list(map(str,range(1,21)))
docchecks={}
for name in ['article','supplement']:
 d=Document(DOC/(name+'.docx'));sec=d.sections[0]
 assert len(d.sections)==1 and abs(sec.page_width.cm-21)<.01 and abs(sec.page_height.cm-29.7)<.01
 for actual,expected in [(sec.top_margin.cm,2),(sec.bottom_margin.cm,2.5),(sec.left_margin.cm,3),(sec.right_margin.cm,1.5)]:assert abs(actual-expected)<.01
 paragraphs=list(d.paragraphs)+[p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs]
 text='\n'.join(p.text for p in paragraphs)
 assert 'HOW TO FORMAT' not in text and 'Given Name Surname' not in text
 for p in paragraphs:
  for r in p.runs:
   if r.text:assert r.font.name=='Times New Roman' and r.font.size.pt==12,(name,r.text)
 with zipfile.ZipFile(DOC/(name+'.docx')) as z,zipfile.ZipFile(BASE/'reference-materials/aitu-article-template.docx') as template:
  for part in load(BASE/'document-build.json')['preserveOnlyParts']:assert z.read(part)==template.read(part)
  omml=z.read('word/document.xml').count(b'<m:oMath>');assert omml==(3 if name=='article' else 0)
 docchecks[name]={'paragraphs':len(paragraphs),'tables':len(d.tables),'nativeOMML':omml,'mainTextFont12TNR':True,'geometry':'A4 top2 bottom2.5 left3 right1.5cm','sha256':sha(DOC/(name+'.docx'))}
validation={'status':'PASS','newControllers':len(new),'controllerStatuses':dict(collections.Counter(r['status'] for r in new)),'pairedRootDiagnostics':12,'rootStatuses':dict(collections.Counter(r['status'] for r in load(BASE/'stage1/root-tables.json'))),'reusedCells':len(reused),'stages':stages,'auditedHashEntriesVerified':hashed,'historicalFilesUnchanged':len(preserved),'frozenProductionSourceHashesEqual':len(sourcechecks),'recompiledClassesByteIdentical':len(rebuild['checks']),'article':{'titleCharacters':len(article['title']),'abstractWords':len(article['abstract'].split()),'abstractCharacters':len(article['abstract']),'referencesInOrder':refOrder},'documents':docchecks,'renderedPages':load(BASE/'document-qa/final2/render.json'),'boundary':'Numerical/structural validation; visual review recorded separately. No new policy runs and no Maven suite in this finalization.'}
save(DOC/'validation.json',validation)
save(BASE/'final-validation.json',validation)
save(BASE/'final-environment.json',{'hardware':load(BASE/'environment-confirmed.json'),'gitCommit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),'gitStatus':subprocess.check_output(['git','status','--porcelain=v1'],cwd=ROOT,text=True),'python':sys.version,'java':rebuild['java'],'jarSha256':sha(ROOT/'target/final-research/final-v2.jar'),'observation':'After campaign on2026-10-28; working tree is dirty, frozen JAR/source manifests define execution identity'})
print(json.dumps({k:v for k,v in validation.items() if k not in ['stages','documents','article','renderedPages']},indent=2))
